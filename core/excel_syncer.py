import os
import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

class ConstructionExcelSyncer:
    """
    Bộ tạo và đồng bộ Bảng tính Excel Báo cáo Tiến độ Công trình chuyên nghiệp.
    Chuẩn mẫu báo cáo Ban Quản Lý Dự Án (PMU) & Tư Vấn Giám Sát.
    Tự động cập nhật công thức sống, định dạng % và màu sắc cảnh báo tiến độ.
    """

    def __init__(self, excel_path: Optional[Path] = None):
        if excel_path is None:
            base_dir = Path(__file__).resolve().parent.parent
            data_dir = base_dir / "data"
            data_dir.mkdir(parents=True, exist_ok=True)
            self.excel_path = data_dir / "Bao_cao_Tien_do_Thi_cong.xlsx"
        else:
            self.excel_path = Path(excel_path)
            self.excel_path.parent.mkdir(parents=True, exist_ok=True)

    def generate_or_update(
        self,
        summary_items: List[Dict[str, Any]],
        piles: List[Dict[str, Any]],
        history_reports: List[Dict[str, Any]],
        project_name: str = "PMU: BĂNG HẠ TẦNG OLP"
    ) -> str:
        """
        Sinh mới hoặc cập nhật toàn bộ file Excel với 3 Sheet chuẩn quản lý dự án.
        """
        wb = openpyxl.Workbook()
        
        # 1. Sheet 1: Dashboard Tổng Hợp Tiến Độ
        ws_dash = wb.active
        ws_dash.title = "Tổng hợp Tiến độ"
        self._build_dashboard_sheet(ws_dash, summary_items, project_name)

        # 2. Sheet 2: Theo Dõi Tim Cọc & Cấu Kiện
        ws_piles = wb.create_sheet(title="Theo dõi Tim Cọc")
        self._build_piles_sheet(ws_piles, piles, project_name)

        # 3. Sheet 3: Lịch Sử Báo Cáo Ca
        ws_history = wb.create_sheet(title="Nhật trình Ca thi công")
        self._build_history_sheet(ws_history, history_reports, project_name)

        # Lưu file
        wb.save(str(self.excel_path))
        return str(self.excel_path)

    def _build_dashboard_sheet(self, ws, items: List[Dict[str, Any]], project_name: str):
        ws.views.sheetView[0].showGridLines = True

        # Styles
        title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
        sub_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
        hdr_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        
        cat_font = Font(name="Segoe UI", size=10, bold=True, color="002060")
        cat_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        
        regular_font = Font(name="Segoe UI", size=9)
        bold_font = Font(name="Segoe UI", size=9, bold=True)

        thin_border_side = Side(style='thin', color='D9D9D9')
        thick_bottom_side = Side(style='medium', color='1F4E79')
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        green_font = Font(name="Segoe UI", size=9, color="006100", bold=True)
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
        yellow_font = Font(name="Segoe UI", size=9, color="9C6500", bold=True)
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_font = Font(name="Segoe UI", size=9, color="9C0006", bold=True)

        # Header Titles
        ws.merge_cells("A1:J1")
        ws["A1"] = f"BÁO CÁO TIẾN ĐỘ THI CÔNG HIỆN TRƯỜNG - {project_name.upper()}"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        ws.merge_cells("A2:J2")
        now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
        ws["A2"] = f"Tự động thu thập và chuẩn hóa dữ liệu bởi: 'Đại ca dữ liệu' | Cập nhật lần cuối: {now_str}"
        ws["A2"].font = sub_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        # Table Headers
        headers = [
            ("STT", 6),
            ("HẠNG MỤC THI CÔNG", 26),
            ("VỊ TRÍ / CẤU KIỆN", 24),
            ("ĐVT", 8),
            ("TỔNG THIẾT KẾ", 16),
            ("CA VỪA QUA", 14),
            ("LŨY KẾ ĐẾN NAY", 16),
            ("TỶ LỆ (%)", 13),
            ("ĐÁNH GIÁ TIẾN ĐỘ", 18),
            ("TRẠNG THÁI / GHI CHÚ", 26)
        ]

        ws.row_dimensions[4].height = 26
        for col_idx, (h_name, width) in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx, value=h_name)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = cell_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        # Group items by Category
        categories_dict: Dict[str, List[Dict[str, Any]]] = {}
        for it in items:
            cat = it.get("category", "Hạng mục chung")
            if cat not in categories_dict:
                categories_dict[cat] = []
            categories_dict[cat].append(it)

        current_row = 5
        stt_counter = 1

        for cat_name, sub_items in categories_dict.items():
            # In dòng Tiêu đề Hạng mục lớn
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
            cat_cell = ws.cell(row=current_row, column=1, value=f"{stt_counter}. {cat_name.upper()}")
            cat_cell.font = cat_font
            cat_cell.fill = cat_fill
            cat_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for c in range(1, 11):
                ws.cell(row=current_row, column=c).border = cell_border
            current_row += 1

            sub_idx = 1
            for sit in sub_items:
                ws.row_dimensions[current_row].height = 20
                
                # STT con
                ws.cell(row=current_row, column=1, value=f"{stt_counter}.{sub_idx}").alignment = Alignment(horizontal="center")
                
                # Hạng mục & Cấu kiện
                ws.cell(row=current_row, column=2, value=cat_name)
                ws.cell(row=current_row, column=3, value=sit.get("sub_item", "-"))
                
                # ĐVT
                ws.cell(row=current_row, column=4, value=sit.get("unit", "Cấu kiện")).alignment = Alignment(horizontal="center")
                
                # Tổng thiết kế (E)
                design_val = float(sit.get("design_qty", 0))
                cell_design = ws.cell(row=current_row, column=5, value=design_val)
                cell_design.number_format = "#,##0"
                cell_design.alignment = Alignment(horizontal="right")
                
                # Ca vừa qua (F)
                shift_val = float(sit.get("shift_qty", 0))
                cell_shift = ws.cell(row=current_row, column=6, value=shift_val)
                cell_shift.number_format = "#,##0"
                cell_shift.alignment = Alignment(horizontal="right")

                # Lũy kế đến nay (G)
                accum_val = float(sit.get("accumulated_qty", 0))
                cell_accum = ws.cell(row=current_row, column=7, value=accum_val)
                cell_accum.number_format = "#,##0"
                cell_accum.alignment = Alignment(horizontal="right")

                # Tỷ lệ % bằng CÔNG THỨC SỐNG: =IF(E{row}>0, G{row}/E{row}, 0)
                cell_pct = ws.cell(row=current_row, column=8, value=f"=IF(E{current_row}>0, G{current_row}/E{current_row}, 0)")
                cell_pct.number_format = "0.0%"
                cell_pct.alignment = Alignment(horizontal="center")

                # Đánh giá tiến độ & Tô màu trực quan
                pct_val = (accum_val / design_val * 100) if design_val > 0 else 0
                cell_eval = ws.cell(row=current_row, column=9)
                cell_eval.alignment = Alignment(horizontal="center")

                if design_val == 0:
                    cell_eval.value = "Đang duy trì"
                elif pct_val >= 80:
                    cell_eval.value = "✅ Gần hoàn thành"
                    cell_pct.fill = green_fill
                    cell_pct.font = green_font
                    cell_eval.fill = green_fill
                    cell_eval.font = green_font
                elif pct_val >= 40:
                    cell_eval.value = "⏳ Đang tăng tốc"
                    cell_pct.fill = yellow_fill
                    cell_pct.font = yellow_font
                    cell_eval.fill = yellow_fill
                    cell_eval.font = yellow_font
                else:
                    cell_eval.value = "⚠️ Khởi động / Chậm"
                    cell_pct.fill = red_fill
                    cell_pct.font = red_font
                    cell_eval.fill = red_fill
                    cell_eval.font = red_font

                # Ghi chú trạng thái
                note_val = sit.get("status_note", "") or "Bình thường"
                ws.cell(row=current_row, column=10, value=note_val)

                # Viền các ô
                for c in range(1, 11):
                    ws.cell(row=current_row, column=c).border = cell_border
                    if ws.cell(row=current_row, column=c).font == Font():
                        ws.cell(row=current_row, column=c).font = regular_font

                current_row += 1
                sub_idx += 1

            stt_counter += 1

    def _build_piles_sheet(self, ws, piles: List[Dict[str, Any]], project_name: str):
        ws.views.sheetView[0].showGridLines = True
        
        hdr_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        thin_border_side = Side(style='thin', color='D9D9D9')
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        regular_font = Font(name="Segoe UI", size=9)

        ws.merge_cells("A1:G1")
        ws["A1"] = f"DANH MỤC TIM CỌC & CẤU KIỆN ĐANG THI CÔNG - {project_name.upper()}"
        ws["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        headers = [
            ("STT", 6),
            ("VỊ TRÍ MỐ / TRỤ", 20),
            ("MÃ TIM CỌC", 18),
            ("TÌNH TRẠNG KỸ THUẬT HIỆN TRƯỜNG", 42),
            ("CA CẬP NHẬT", 14),
            ("NGÀY BÁO CÁO", 14),
            ("NHÀ THẦU THI CÔNG", 22)
        ]

        ws.row_dimensions[3].height = 24
        for col_idx, (h_name, width) in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h_name)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        row_idx = 4
        for idx, p in enumerate(piles, 1):
            ws.row_dimensions[row_idx].height = 20
            ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=p.get("location", "-"))
            ws.cell(row=row_idx, column=3, value=p.get("pile_id", "-")).font = Font(name="Segoe UI", size=9, bold=True)
            
            st_cell = ws.cell(row=row_idx, column=4, value=p.get("status", "-"))
            if "khoan" in p.get("status", "").lower():
                st_cell.font = Font(name="Segoe UI", size=9, color="B25900", bold=True)
            elif "bê tông" in p.get("status", "").lower():
                st_cell.font = Font(name="Segoe UI", size=9, color="008000", bold=True)
            elif "lồng thép" in p.get("status", "").lower():
                st_cell.font = Font(name="Segoe UI", size=9, color="002060", bold=True)
                
            ws.cell(row=row_idx, column=5, value=p.get("shift_name", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=6, value=p.get("report_date", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=7, value=p.get("contractor_name", "-"))

            for c in range(1, 8):
                ws.cell(row=row_idx, column=c).border = cell_border
                if not ws.cell(row=row_idx, column=c).font.name:
                    ws.cell(row=row_idx, column=c).font = regular_font
            row_idx += 1

    def _build_history_sheet(self, ws, history: List[Dict[str, Any]], project_name: str):
        ws.views.sheetView[0].showGridLines = True
        
        hdr_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        thin_border_side = Side(style='thin', color='D9D9D9')
        cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        regular_font = Font(name="Segoe UI", size=9)

        ws.merge_cells("A1:F1")
        ws["A1"] = f"NHẬT TRÌNH BÁO CÁO CA THI CÔNG - {project_name.upper()}"
        ws["A1"].font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        headers = [
            ("MÃ BC", 8),
            ("NGÀY BÁO CÁO", 15),
            ("CA THI CÔNG", 15),
            ("NHÀ THẦU", 22),
            ("NGƯỜI BÁO CÁO", 18),
            ("SỐ HẠNG MỤC NẠP", 18)
        ]

        ws.row_dimensions[3].height = 24
        for col_idx, (h_name, width) in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h_name)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        row_idx = 4
        for h in history:
            ws.row_dimensions[row_idx].height = 20
            ws.cell(row=row_idx, column=1, value=f"BC#{h.get('id', 0)}").alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=h.get("report_date", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=3, value=h.get("shift_name", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=4, value=h.get("contractor_name", "-"))
            ws.cell(row=row_idx, column=5, value=h.get("reporter_name", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=6, value=h.get("item_count", 0)).alignment = Alignment(horizontal="right")

            for c in range(1, 7):
                ws.cell(row=row_idx, column=c).border = cell_border
                ws.cell(row=row_idx, column=c).font = regular_font
            row_idx += 1

    def export_search_results(
        self,
        items: List[Dict[str, Any]],
        piles: Optional[List[Dict[str, Any]]] = None,
        filter_summary: str = "Tất cả dữ liệu",
        target_path: Optional[Path] = None
    ) -> str:
        """Xuất kết quả tìm kiếm/lọc dữ liệu thành file Excel riêng biệt để báo cáo."""
        out_path = target_path or (self.excel_path.parent / "Ket_qua_Loc_Du_lieu.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dữ liệu lọc"
        ws.views.sheetView[0].showGridLines = True

        title_font = Font(name="Segoe UI", size=13, bold=True, color="1F4E79")
        sub_font = Font(name="Segoe UI", size=9, italic=True, color="595959")
        hdr_font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
        hdr_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        regular_font = Font(name="Segoe UI", size=9)
        bold_font = Font(name="Segoe UI", size=9, bold=True)

        ws.merge_cells("A1:K1")
        ws["A1"] = f"BẢNG KẾT QUẢ TRA CỨU & LỌC DỮ LIỆU CÔNG TRƯỜNG"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 26

        ws.merge_cells("A2:K2")
        ws["A2"] = f"Điều kiện lọc: {filter_summary} | Xuất lúc: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        ws["A2"].font = sub_font
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 18

        headers = [
            ("STT", 6),
            ("DỰ ÁN / NHÓM ZALO", 22),
            ("NGÀY", 12),
            ("CA THI CÔNG", 14),
            ("NHÀ THẦU", 18),
            ("HẠNG MỤC", 24),
            ("VỊ TRÍ / CẤU KIỆN", 22),
            ("ĐVT", 8),
            ("CA NÀY", 12),
            ("LŨY KẾ", 12),
            ("TỔNG TK", 12),
            ("TIẾN ĐỘ (%)", 14),
            ("GHI CHÚ / TIM CỌC", 28)
        ]

        ws.row_dimensions[3].height = 22
        for col_idx, (h_name, width) in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=h_name)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cell_border
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        row_idx = 4
        for idx, it in enumerate(items, 1):
            ws.row_dimensions[row_idx].height = 20
            ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=2, value=it.get("project_name", "-"))
            ws.cell(row=row_idx, column=3, value=it.get("report_date", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=4, value=it.get("shift_name", "-")).alignment = Alignment(horizontal="center")
            ws.cell(row=row_idx, column=5, value=it.get("contractor_name", "-"))
            ws.cell(row=row_idx, column=6, value=it.get("category", "-"))
            ws.cell(row=row_idx, column=7, value=it.get("sub_item", "-"))
            ws.cell(row=row_idx, column=8, value=it.get("unit", "Cấu kiện")).alignment = Alignment(horizontal="center")
            
            c_shift = ws.cell(row=row_idx, column=9, value=it.get("shift_qty", 0))
            c_shift.number_format = '#,##0.00' if isinstance(it.get("shift_qty"), float) and not it.get("shift_qty").is_integer() else '#,##0'
            c_shift.alignment = Alignment(horizontal="right")

            c_acc = ws.cell(row=row_idx, column=10, value=it.get("accumulated_qty", 0))
            c_acc.number_format = '#,##0.00' if isinstance(it.get("accumulated_qty"), float) and not it.get("accumulated_qty").is_integer() else '#,##0'
            c_acc.alignment = Alignment(horizontal="right")

            c_des = ws.cell(row=row_idx, column=11, value=it.get("design_qty", 0))
            c_des.number_format = '#,##0.00' if isinstance(it.get("design_qty"), float) and not it.get("design_qty").is_integer() else '#,##0'
            c_des.alignment = Alignment(horizontal="right")

            rate = it.get("completion_rate", 0)
            c_pct = ws.cell(row=row_idx, column=12, value=rate / 100.0)
            c_pct.number_format = '0.0%'
            c_pct.alignment = Alignment(horizontal="right")

            ws.cell(row=row_idx, column=13, value=it.get("status_note", "-"))

            for c in range(1, 14):
                ws.cell(row=row_idx, column=c).border = cell_border
                ws.cell(row=row_idx, column=c).font = regular_font
            row_idx += 1

        if piles:
            ws_p = wb.create_sheet(title="Tim cọc lọc")
            ws_p.views.sheetView[0].showGridLines = True
            p_headers = [("STT", 6), ("DỰ ÁN / NHÓM ZALO", 22), ("NGÀY", 12), ("CA", 14), ("NHÀ THẦU", 18), ("VỊ TRÍ / MỐ TRỤ", 18), ("MÃ HIỆU CỌC", 16), ("TRẠNG THÁI HIỆN TRƯỜNG", 32)]
            for col_idx, (h_name, width) in enumerate(p_headers, 1):
                cell = ws_p.cell(row=1, column=col_idx, value=h_name)
                cell.font = hdr_font
                cell.fill = hdr_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = cell_border
                ws_p.column_dimensions[get_column_letter(col_idx)].width = width
            for p_idx, p in enumerate(piles, 1):
                r = p_idx + 1
                ws_p.cell(row=r, column=1, value=p_idx).alignment = Alignment(horizontal="center")
                ws_p.cell(row=r, column=2, value=p.get("project_name", "-"))
                ws_p.cell(row=r, column=3, value=p.get("report_date", "-")).alignment = Alignment(horizontal="center")
                ws_p.cell(row=r, column=4, value=p.get("shift_name", "-")).alignment = Alignment(horizontal="center")
                ws_p.cell(row=r, column=5, value=p.get("contractor_name", "-"))
                ws_p.cell(row=r, column=6, value=p.get("location", "-"))
                ws_p.cell(row=r, column=7, value=p.get("pile_id", "-")).font = bold_font
                ws_p.cell(row=r, column=8, value=p.get("status", "-"))
                for c in range(1, 9):
                    ws_p.cell(row=r, column=c).border = cell_border
                    ws_p.cell(row=r, column=c).font = regular_font

        wb.save(str(out_path))
        return str(out_path)
